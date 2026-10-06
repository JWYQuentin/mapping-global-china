"""Add Phase 1d (non-site data: agreements, partnerships, projects) sheets to the Phase 1 workbook.
Run from the repo root, in order after the earlier build scripts (each script adds sheets once).
All values were collected on 2026-10-05 from the URLs recorded in each row."""
import sys
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

SRC = sys.argv[1] if len(sys.argv) > 1 else "data/processed/Mapping_Global_China_Phase1_Sources.xlsx"
OUT = sys.argv[2] if len(sys.argv) > 2 else "data/processed/Mapping_Global_China_Phase1_Sources.xlsx"
F = "Arial"
HDR = PatternFill("solid", fgColor="1F3A5F"); HFONT = Font(name=F, bold=True, color="FFFFFF")
BODY = Font(name=F, size=10); LINK = Font(name=F, size=10, color="0563C1", underline="single")
FILL = PatternFill("solid", fgColor="FFF2CC")
wb = load_workbook(SRC)

def sheet(title, headers, rows, widths, link_cols=(), fill_cols=(), freeze="C2"):
    ws = wb.create_sheet(title); ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c); cell.font, cell.fill = HFONT, HDR
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for r in rows: ws.append(list(r))
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = BODY; cell.alignment = Alignment(wrap_text=True, vertical="top")
            if cell.column in link_cols and isinstance(cell.value, str) and cell.value.startswith("http"):
                cell.hyperlink = cell.value; cell.font = LINK
            if cell.column in fill_cols: cell.fill = FILL
    for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = freeze; ws.auto_filter.ref = ws.dimensions; ws.row_dimensions[1].height = 32

TDB = "https://treaty.mfa.gov.cn/web/detail1.jsp?objid="
# id, title, type, field, signed, in force, matched keyword | partner (my reading), category (my coding)
TR = [
 ("1531876945033","中华人民共和国国家文物局与澳大利亚环境、水、遗产和艺术部关于文物保护的谅解备忘录","双边","文化","2009.10.30","2009.10.30","文物","Australia","Heritage cooperation"),
 ("1531876933822","中华人民共和国国家文物局与埃塞俄比亚联邦民主共和国文化遗产研究与保护局关于合作保护文化遗产谅解备忘录","双边","文化","2009.02.25","2009.02.25","文物","Ethiopia","Heritage cooperation"),
 ("1531876925884","中华人民共和国国家文物局和智利共和国国家遗迹委员会关于开展文化遗产领域交流合作的协议","双边","文化","2008.04.14","2008.04.14","文物","Chile","Heritage cooperation"),
 ("1531876925835","中华人民共和国政府和智利共和国政府关于防止盗窃、盗掘和非法进出境文物的协定","双边","其他","2008.04.14","2008.07.01","文物","Chile","Illicit trafficking"),
 ("1531876924395","中华人民共和国国家文物局和希腊共和国文化部关于防止盗窃、盗掘和非法进出境文物的谅解备忘录","双边","文化","2008.02.26","","文物","Greece","Illicit trafficking"),
 ("1531876918053","中华人民共和国国家文物局与斯里兰卡民主社会主义共和国国家遗产部关于合作保护文化遗产的协议","双边","文化","2008.01.31","2008.01.31","文物","Sri Lanka","Heritage cooperation"),
 ("1531876904849","中华人民共和国国家文物局与伊朗伊斯兰共和国文化遗产手工业旅游局关于合作保护文化遗产谅解备忘录","双边","文化","2007.09.07","2007.09.07","文物","Iran","Heritage cooperation"),
 ("1531876903025","中华人民共和国政府和菲律宾共和国政府关于防止盗窃、盗掘和非法进出境文物的协定","双边","其他","2007.01.15","","文物","Philippines","Illicit trafficking"),
 ("1531876897634","中华人民共和国政府和意大利共和国政府关于防止盗窃、盗掘和非法进出境文物的协定","双边","其他","2006.01.20","","文物","Italy","Illicit trafficking"),
 ("1531876893427","中华人民共和国国家文物局与越南社会主义共和国文化通讯部文化遗产局关于文化遗产保护合作谅解备忘录","双边","文化","2006.10.11","","文物","Viet Nam","Heritage cooperation"),
 ("1531876893282","中华人民共和国国家文物局与印度共和国考古局关于合作保护文化遗产谅解备忘录","双边","文化","2006.11.21","2006.11.21","文物","India","Heritage cooperation"),
 ("1531876893181","中华人民共和国政府和印度共和国政府关于防止盗窃、盗掘和非法进出境文物的协定","双边","其他","2006.11.21","","文物","India","Illicit trafficking"),
 ("1531876888674","中华人民共和国国家文物局与联合国教育、科学、文化组织关于信托基金项目的行动计划","双边","文化","","","文物","UNESCO","Funds-in-trust"),
 ("1531876878468","中华人民共和国国家文物局和大韩民国文化财厅关于开展文化遗产方面交流合作的协议","双边","文化","2005.11.03","2005.11.03","文物","Republic of Korea","Heritage cooperation"),
 ("1531876565956","中国社会科学院和瑞典皇家文学、历史、文物学院学术交流协议","双边","文化","1981.10.07","1981.10.07","文物","Sweden","Academic exchange"),
 ("1531876513690","中华人民共和国政府和澳大利亚政府关于《中华人民共和国出土文物展览》协定","双边","文化","1976.06.23","1976.06.23","文物","Australia","Exhibition (1970s)"),
 ("1531876513639","中华人民共和国政府和菲律宾共和国政府关于《中华人民共和国出土文物展览》协定","双边","文化","1976.08.05","1976.08.05","文物","Philippines","Exhibition (1970s)"),
 ("1531876494252","中华人民共和国政府和墨西哥合众国政府关于《中华人民共和国出土文物展览》协定","双边","文化","1974.05.24","1974.05.24","文物","Mexico","Exhibition (1970s)"),
 ("1531876494151","中华人民共和国政府和荷兰王国政府关于《中华人民共和国出土文物展览》协定","双边","文化","1974.10.28","1974.10.28","文物","Netherlands","Exhibition (1970s)"),
 ("1531876494044","中华人民共和国政府和加拿大政府关于《中华人民共和国出土文物展览》协定","双边","文化","1974.03.15","1974.03.15","文物","Canada","Exhibition (1970s)"),
 ("1531876493991","中华人民共和国政府和比利时王国政府关于《中华人民共和国出土文物展览》协定","双边","文化","1974.12.20","1974.12.20","文物","Belgium","Exhibition (1970s)"),
 ("1531876486014","中华人民共和国政府和瑞典王国政府关于《中华人民共和国出土文物展览》协定","双边","文化","1973.11.09","1973.12.10","文物","Sweden","Exhibition (1970s)"),
 ("1531876485962","中华人民共和国政府和奥地利共和国政府关于《中华人民共和国出土文物展览》协定","双边","文化","1973.10.30","1974.01.31","文物","Austria","Exhibition (1970s)"),
 ("1531876485908","中华人民共和国政府和大不列颠及北爱尔兰联合王国政府关于《中华人民共和国出土文物展览》协定","双边","文化","1973.07.02","1973.07.02","文物","United Kingdom","Exhibition (1970s)"),
 ("1531876485802","中华人民共和国政府和法兰西共和国政府关于《中华人民共和国出土文物展览》协定","双边","文化","1973.03.23","1973.03.23","文物","France","Exhibition (1970s)"),
 ("1531876485752","中华人民共和国政府和罗马尼亚社会主义共和国政府关于《中华人民共和国出土文物展览》协定","双边","文化","1973.11.24","1973.11.24","文物","Romania","Exhibition (1970s)"),
 ("1531876070902","国际统一私法协会关于被盗或者非法出口文物的公约","多边","文化","","1998.07.01","文物","Multilateral (UNIDROIT)","Multilateral convention"),
 ("1531876854962","中华人民共和国文化部和匈牙利共和国民族文化遗产部二〇〇三至二〇〇五年文化合作计划","双边","文化","2003.01.27","2003.01.27","文化遗产","Hungary","Cultural cooperation plan"),
 ("1531876077415","保护非物质文化遗产公约","多边","文化","","2006.04.20","文化遗产","Multilateral (UNESCO)","Multilateral convention"),
 ("1605167857701","中华人民共和国政府和秘鲁共和国政府关于对秘鲁国家考古、人类、 历史学博物馆附属设施项目开展可行性研究的换函","双边","文化","2016.11.21","","考古","Peru","Aid project (museum)"),
 ("1531876944132","中华人民共和国政府和美利坚合众国政府对旧石器时代到唐末的归类考古材料以及至少250年以上的古迹雕塑和壁上艺术实施进口限制的谅解备忘录","双边","投资贸易","2009.01.14","2009.01.14","考古","United States","Illicit trafficking"),
]
rows = sorted(([r[7], r[8], r[1], r[2], r[3], r[4], r[5], r[6], TDB + r[0], ""] for r in TR), key=lambda x: (x[5] or "9999", x[0]))
sheet("1d Agreements (MFA DB)",
      ["Partner (my reading)", "Category (my coding)", "Title (original, as listed)", "Type (类别)", "Field (领域)",
       "Signed by China (我国签署时间)", "In force (条约生效时间)", "Matched keyword", "Record", "Your notes"],
      rows, [22, 24, 70, 10, 10, 16, 16, 14, 46, 24], link_cols=(9,), fill_cols=(10,))

INV = [
 ("Agreements", "China Ministry of Foreign Affairs Treaty Database (条约数据库)", "Chinese government", "https://treaty.mfa.gov.cn/web/index.jsp",
  "Bilateral and multilateral treaties and some cooperation documents. 31 records match 文物, 文化遗产, 考古, 世界遗产 or 非物质文化遗产: 6 on illicit trafficking, 8 heritage-cooperation memoranda, 11 exhibition agreements from 1973-76.",
  "Chinese", "Highly structured: title, type, field, date signed, date in force, record ID, full text.",
  "Keyword search by web form (tested 2026-10-05); results can be collected page by page.", "Yes: signing dates from 1973.",
  "Incomplete after 2009: it holds 5 country agreements on illicit trafficking plus the US memorandum, while officials cite 26-27. Titles shift from 'exhibition of unearthed relics' (1970s) to 'preventing theft and illicit export' and 'cooperating to protect heritage' (2005-09).", "High for what it holds.", ""),
 ("Agreements", "Official counts of bilateral agreements on illicit trafficking of cultural property", "Chinese media citing officials", "https://m.bjnews.com.cn/detail/1699965728129667.html",
  "26 countries as of November 2023 (Beijing News); 27 as of September 2026 (Yangcheng Evening News).", "Chinese",
  "Counts only. No public list of the countries found.", "Manual; ask NCHA or assemble from news.", "Counts are dated.",
  "Framed as recovering lost relics.", "Medium (media).", ""),
 ("Agreements", "US State Department: cultural property agreements and import restrictions", "Foreign government", "https://2012-2025.eca.state.gov/cultural-heritage-center/cultural-property/current-agreements-and-import-restrictions.html",
  "41 countries with US agreements or restrictions. China: memorandum of 14 January 2009, extended 2014, 2019 and 2024.", "English",
  "Structured list by country with Federal Register notices and designated lists.", "Web page (archived copy covering 2012 to May 2025) and Federal Register.", "Yes: each extension dated.",
  "Counterparty view of the same agreement; Chinese press reports more than 700 items returned in over 20 batches since 2009.", "High.", ""),
 ("Legal frameworks", "UNESCO Database of National Cultural Heritage Laws (China page)", "UNESCO", "https://whc.unesco.org/en/statesparties/CN/laws/",
  "19 Chinese legal documents, from 1928 regulations to the Law on Protection of Cultural Relics (1982, 2002, 2007); includes bilateral agreements with Egypt (2010) and Switzerland (2013).", "Chinese, English, French",
  "Structured table: document, year, theme, category, language, with PDFs.", "Web page with PDF links; same structure for every country.", "Yes: dated versions of laws.",
  "Neutral. Lets you compare host-country heritage laws in the same format.", "High; may lag recent amendments.", ""),
 ("UNESCO meetings", "World Heritage Committee sessions archive and decisions database", "UNESCO", "https://whc.unesco.org/en/sessions/",
  "Every Committee, Bureau and General Assembly session since 1977 with documents, decisions and records. Sessions held in China: 28th (Suzhou, 2004) and 44th (Fuzhou, 2021).", "English, French",
  "Highly structured by session code; decisions have stable IDs.", "Web pages; UNESCO's terms on scraping apply (see Data Access sheet).", "Yes: full dated series.",
  "Official wording of decisions; useful for when terms such as 'Silk Roads' enter decisions.", "High.", ""),
 ("Partnerships and funding", "UNESCO Core Data Portal", "UNESCO", "https://core.unesco.org/en/project/580RAF4001",
  "Project-level data by donor and sector. Example: China Funds-in-Trust (Quanzhou) for Capacity-Building on World Heritage in Africa, USD 1,000,000, 11 March 2025 to 10 March 2028.", "English",
  "Highly structured; XLS download.", "Download or page per project ID.", "Yes: project dates and budgets.",
  "Shows a Chinese city (Quanzhou, a 2021 World Heritage site) as named donor for African nominations.", "High.", ""),
 ("Partnerships and funding", "UNESCO World Heritage Centre partner page: Government of China", "UNESCO", "https://whc.unesco.org/en/partners/503",
  "China-funded World Heritage work in Africa since 2014: nomination support in Chad, Cameroon, Nigeria, Benin and Togo; capacity building in Uganda, Congo, Ethiopia, Gabon, Mali, Côte d'Ivoire, Senegal and Tanzania; 2019 UNESCO-Africa-China Forum.", "English",
  "One page per partner with linked news and events. No amounts given.", "Web page.", "Linked news items are dated.",
  "China as funder of other countries' nominations, not of 'Chinese' sites.", "High.", ""),
 ("Training", "UNESCO category 2 centres hosted by China", "UNESCO / Chinese institutions", "https://whc.unesco.org/en/activities/676/",
  "World Heritage centres: WHITR-AP (training and research, Shanghai) and HIST (space technologies). Intangible heritage: CRIHAP (training centre for Asia-Pacific, set up 2012; a 2019 workshop drew beneficiaries from Cambodia, Kazakhstan, Mongolia, Myanmar, Nepal and others).", "English, Chinese",
  "Each centre publishes activity news and annual reports; no shared database.", "Manual, or scrape each centre's news.", "Yes: annual reports.",
  "China as a regional training provider.", "High.", ""),
 ("Partnerships and funding", "Alliance for Cultural Heritage in Asia (ACHA) and the Asian Fund for Cultural Heritage Conservation", "China-led intergovernmental organisation; fund run under the China Foundation for Cultural Heritage Conservation", "http://www.news.cn/world/20251128/1965e348e6f14010aca3d1de668795ef/c.html",
  "ACHA became an intergovernmental organisation on 19 June 2024 (20 members, observers and partners at that date). Fund rounds: first 15 projects, RMB 4.6364 million, 8 countries (announced 2024); 15 more for 2024-2025 (announced November 2025, Chongqing).",
  "Chinese, English", "Project lists announced through state media; no public project database found.", "Manual from news; ask ACHA for lists.", "Yes: rounds dated 2022, 2024, 2025.",
  "'Asian' framing: projects in Armenia, Azerbaijan, Cambodia, UAE, Sri Lanka, Syria, Uzbekistan. New committee on 'Austronesian peoples and maritime civilisation'.", "Medium-high (state media).", ""),
 ("Projects", "Belt and Road Forum deliverables lists (2017, 2019, 2023)", "Chinese government", "https://www.mfa.gov.cn/zyxw/202310/t20231018_11163412.shtml",
  "Numbered lists of agreed projects. The October 2023 list includes the Maritime Silk Road transnational nomination, a China-Saudi joint excavation at Serrin port, and an archaeological park at Nateshwar, Bangladesh.", "Chinese, English",
  "Semi-structured numbered lists, one document per forum.", "Web page; text extraction.", "Yes: three dated lists.",
  "Heritage listed as a Belt and Road deliverable alongside infrastructure.", "High for intent; not for delivery.", ""),
 ("Projects", "State-media overviews of aid restoration", "Chinese media", "https://news.gmw.cn/2025-01/02/content_37772689.htm",
  "Guangming Daily (2 January 2025): restoration aid in Cambodia, Mongolia, Uzbekistan (Khiva), Nepal, Myanmar (Bagan) and Kyrgyzstan.", "Chinese",
  "Narrative articles.", "Manual.", "Dated articles.",
  "'One flower alone is not spring; a hundred flowers in bloom fill the garden' (civilisational diversity).", "Medium.", ""),
 ("Projects", "Chinese companies and private donors in heritage work abroad", "Not found", "",
  "No source found that lists Chinese companies restoring or funding heritage sites abroad.", "", "None found.",
  "Places to look next: implementing-agency field in AidData (1c sheet); Ministry of Commerce aid tenders; donor lists of the China Foundation for Cultural Heritage Conservation.", "Unknown.", "Not sampled.", "Unknown.", ""),
 ("Projects", "Regional summit outcome lists (China-Central Asia Summit, FOCAC action plans)", "Chinese government", "https://uz.china-embassy.gov.cn/zwgxzxdt/202506/t20250618_11652102.htm",
  "Outcome lists and action plans that may contain heritage clauses.", "Chinese, English",
  "Semi-structured lists.", "Not yet read; test in a later pass.", "Yes: one per summit.", "Not sampled.", "High for intent.", ""),
]
sheet("1d Non-site Data Sources",
      ["Data type", "Source", "Owner type", "URL (example page)", "Coverage", "Language", "How systematic",
       "Access method (preliminary, for Phase 2)", "Historical versions", "Framing notes", "Reliability", "Your notes"],
      INV, [18, 36, 24, 40, 56, 12, 40, 44, 26, 52, 20, 24], link_cols=(4,), fill_cols=(12,))

EV = [
 ("2009-01-14", "Agreement", "China; United States", "Memorandum restricting US imports of Chinese archaeological material; extended 2014, 2019, 2024.", "", "https://2012-2025.eca.state.gov/cultural-heritage-center/cultural-property/current-agreements-and-import-restrictions.html", "Retrieved by web-reading tool; check"),
 ("2022", "Fund project", "China; Cambodia", "'Angkor Guardians': 2,000 care packs for children at 6 primary schools near Angkor, delivered with APSARA and the Chinese Academy of Cultural Heritage.",
  "由亚洲文化遗产保护基金和中国爱心网友提供的2000个“爱心包”已于2022年11月运送到柬埔寨", "http://www.news.cn/culture/20230320/1b237d8ebd314e15be1fb95123e8be6e/c.html", "Copied from page"),
 ("2022", "Fund project", "China; Afghanistan", "'Bamiyan Guardians': training for young heritage staff and residents of the Bamiyan site.",
  "完成了15次“馕课”教学活动，针对性培训遗产保护区和缓冲区原住民300人以上", "http://www.news.cn/culture/20230320/1b237d8ebd314e15be1fb95123e8be6e/c.html", "Copied from page"),
 ("2023-10-18", "Belt and Road deliverable", "China; Saudi Arabia; Bangladesh", "NCHA to pursue the Maritime Silk Road transnational joint nomination, a China-Saudi excavation at Serrin port and an archaeological park at Nateshwar.",
  "中国国家文物局将开展海上丝绸之路跨国联合申遗工作，举办“海上丝绸之路文化遗产保护圆桌会议”、五湖四海-“一带一路”文物考古合作展，并形成“一带一路”文化遗产保护成果图集。实施中国—沙特塞林港联合考古项目、援孟加拉国耐特什瓦考古遗址公园项目。", "https://www.mfa.gov.cn/zyxw/202310/t20231018_11163412.shtml", "Copied from page"),
 ("2023-10-18", "Belt and Road deliverable", "China", "White paper on Belt and Road cultural heritage protection.", "发布《“一带一路”文化遗产保护白皮书》。", "https://www.mfa.gov.cn/zyxw/202310/t20231018_11163412.shtml", "Copied from page"),
 ("2023-11-14", "Official count", "China", "26 countries have signed agreements with China on preventing illicit import and export of cultural property.", "中国已与26个国家签署防止文物非法进出境政府间协定", "https://m.bjnews.com.cn/detail/1699965728129667.html", "Retrieved by web-reading tool; check"),
 ("2024-06-19", "Organisation", "ACHA members", "ACHA becomes an intergovernmental organisation; Uzbekistan and Maldives join, Honduras becomes first partner country; first 15 fund projects approved (RMB 4.6364 million, 8 countries).",
  "大会吸纳乌兹别克斯坦、马尔代夫成为联盟成员国，洪都拉斯成为联盟首个伙伴国，联盟成员国、观察员国、伙伴国总数达到20个。", "https://cn.chinadaily.com.cn/a/202406/20/WS66738c64a3107cd55d267adc.html", "Retrieved by web-reading tool; check"),
 ("2025-03-11", "UNESCO funds-in-trust", "China (Quanzhou); Africa", "China Funds-in-Trust (Quanzhou) for Capacity-Building on World Heritage in Africa, USD 1,000,000, to March 2028.", "", "https://core.unesco.org/en/project/580RAF4001", "Retrieved by web-reading tool; check"),
 ("2025-11-27", "Organisation", "ACHA members; Egypt", "ACHA second general assembly, Chongqing: 15 fund projects for 2024-2025; host-country agreement with China; NCHA-Egypt joint statement and World Heritage cooperation agreement.",
  "本届大会发布了2024—2025年度亚洲文化遗产保护基金资助项目，涉及亚美尼亚、阿塞拜疆、柬埔寨、阿联酋等国家的15个项目。", "http://www.news.cn/world/20251128/1965e348e6f14010aca3d1de668795ef/c.html", "Copied from page"),
 ("2025-11-27", "Fund projects", "China; Uzbekistan; Sri Lanka; Syria", "Named projects: Uzbekistan-China excavation at Chinar-tepe; excavation at Papalugala and study of Mantai port finds, Sri Lanka; Latakia Museum exhibition upgrade.",
  "乌兹别克斯坦—中国联合考古发掘契纳尔特佩遗址、斯里兰卡帕巴鲁伽拉考古发掘与曼泰港遗址文物研究、拉塔基亚博物馆展览提升等15个项目将得到亚洲文化遗产保护基金资助。", "http://www.news.cn/world/20251128/1965e348e6f14010aca3d1de668795ef/c.html", "Copied from page"),
 ("2026-09-19", "Official count", "China; United States", "27 countries have bilateral agreements with China; over 700 items returned from the US in more than 20 batches since 2009.",
  "中美两国签署防止中国文物非法入境美国的政府间谅解备忘录以来，务实合作成果丰硕，实现超过20批次共700余件（套）流失美国的文物艺术品回归中国。", "https://m.gmw.cn/2026-09/19/content_1304565539.htm", "Retrieved by web-reading tool; check"),
]
sheet("1d Programmes & Outcomes",
      ["Date", "Type", "Parties", "What (my summary)", "Original wording", "Source", "Verification status"],
      EV, [12, 22, 28, 60, 64, 44, 26], link_cols=(6,), fill_cols=(7,), freeze="B2")

ws = wb["README"]
ws["A1"] = "Mapping Global China - Phase 1 source inventory (1a UNESCO, 1b Chinese sources, 1c host-country and diaspora, 1d non-site data)"
r = ws.max_row + 2
add = [
 ("Phase 1d (added 2026-10-05)", None),
 ("Non-site data sources surveyed", "=COUNTA('1d Non-site Data Sources'!B2:B200)"),
 ("Agreements found in China's treaty database", "=COUNTA('1d Agreements (MFA DB)'!C2:C200)"),
 ("  of which on illicit trafficking", "=COUNTIF('1d Agreements (MFA DB)'!B2:B200,\"Illicit trafficking\")"),
 ("  of which heritage cooperation", "=COUNTIF('1d Agreements (MFA DB)'!B2:B200,\"Heritage cooperation\")"),
 ("Programmes and outcomes recorded", "=COUNTA('1d Programmes & Outcomes'!A2:A200)"),
 ("1d Non-site Data Sources", "Inventory of sources for agreements, legal frameworks, UNESCO meetings, partnerships, training and projects."),
 ("1d Agreements (MFA DB)", "All records matching five heritage keywords in the Foreign Ministry treaty database, copied from the database. Partner and category columns are my coding."),
 ("1d Programmes & Outcomes", "Dated seed records of funds, organisations and deliverables, with original wording where copied from the page."),
]
for i, (a, b) in enumerate(add):
    ws.cell(row=r + i, column=1, value=a).font = Font(name=F, size=11 if i == 0 else 10, bold=(i == 0))
    if b is not None:
        c = ws.cell(row=r + i, column=2, value=b); c.font = Font(name=F, size=10); c.alignment = Alignment(wrap_text=True, vertical="top")
wb.save(OUT); print("saved", OUT, len(INV), len(TR), len(EV))
