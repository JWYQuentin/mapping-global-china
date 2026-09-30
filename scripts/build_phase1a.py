"""Build the Phase 1a UNESCO inventory workbook.

Run from the repo root: python scripts/build_phase1a.py
Input: data/raw/unesco_whc_feed_china_extract_2026-09-30.json (China records from
https://whc.unesco.org/en/list/xml/, extracted 2026-09-30). Other values were read from
UNESCO web pages on 2026-09-30 and are recorded inline below.
"""
import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ASOF = "2026-09-30"
WHC = "https://whc.unesco.org"

ZH = {"144":"基尔瓦基斯瓦尼遗址和松戈马拉遗址","208":"巴米扬山谷的文化景观和考古遗迹","340":"凯奥拉德奥国家公园","365":"卡米国家遗址纪念地","437":"泰山","438":"长城","439":"明清故宫（北京故宫、沈阳故宫）","440":"莫高窟","441":"秦始皇陵及兵马俑坑","449":"周口店北京人遗址","502":"维甘历史古城","547":"黄山","559":"德罗特宁霍尔摩皇宫","637":"九寨沟风景名胜区","638":"黄龙风景名胜区","640":"武陵源风景名胜区","660":"法隆寺地区的佛教古迹","677":"菲律宾的巴洛克教堂","688":"古京都遗址（京都、宇治和大津城）","703":"承德避暑山庄及其周围寺庙","704":"曲阜孔庙、孔林和孔府","705":"武当山古建筑群","707":"拉萨布达拉宫历史建筑群","778":"庐山国家公园","779":"峨眉山—乐山大佛","811":"丽江古城","812":"平遥古城","813":"苏州古典园林","870":"古奈良的历史遗迹","880":"北京皇家园林-颐和园","881":"北京皇家祭坛—天坛","911":"武夷山","912":"大足石刻","948":"会安古镇","972":"琉球王国时期的遗迹","1001":"青城山—都江堰","1002":"皖南古村落－西递、宏村（2000年）","1003":"龙门石窟","1004":"明清皇家陵寝","1012":"基纳巴卢山公园","1039":"云岗石窟","1083":"云南三江并流保护区","1091":"高句丽古墓群","1110":"澳门历史城区","1112":"开平碉楼与村落","1113":"福建土楼","1114":"殷墟","1135":"高句丽王城、王陵及贵族墓葬","1142":"纪伊山地的圣地与参拜道","1213":"四川大熊猫栖息地","1246":"石见银山遗迹及其文化景观","1248":"中国南方喀斯特","1279":"五台山","1292":"三清山国家公园","1305":"登封 “天地之中”历史古迹","1328":"河内升龙皇城","1334":"杭州西湖文化景观","1335":"中国丹霞","1346":"大不里士的集市区","1388":"澄江化石遗址","1389":"元上都遗址","1474":"土司遗址","1477":"百济遗址区","1498":"韩国新儒学书院","1508":"左江花山岩画文化景观","1509":"湖北神农架","1537":"卡尔哈特古城","1559":"梵净山","1561":"泉州：宋元中国的世界海洋商贸中心","1574":"奄美大岛、德之岛、冲绳岛北部及西表岛","1592":"良渚古城遗址","1606":"中国黄（渤）海候鸟栖息地","1610":"沙哇伦多的翁比林煤矿遗产","1627":"古骨咄的文化遗产地","1638":"巴丹吉林沙漠——沙山湖泊群","1665":"普洱景迈山古茶林文化景观","1675":"丝绸之路：扎拉夫尚-卡拉库姆廊道","1714":"北京中轴线：中国理想都城秩序的杰作","1736":"西夏陵","1757":"飞鸟与藤原古都","1759":"匈奴贵族墓葬群","1765":"景德镇手工瓷业遗存"}

# China sites: links stated in UNESCO's own short description (EN feed)
CN_FLAGS = {
 "1442": ("Silk Roads", "This property is a 5,000 km section of the extensive Silk Roads network, stretching from Chang'an/Luoyang, the central capital of China in the Han and Tang dynasties, to the Zhetysu region of Central Asia."),
 "440": ("Silk Roads", "Situated at a strategic point along the Silk Route, at the crossroads of trade as well as religious, cultural and intellectual influences..."),
 "1736": ("Silk Roads", "Positioned along the Silk Road, it became a multicultural civilization modelled on Chinese imperial traditions, with Buddhism at its core."),
 "1561": ("Maritime trade", "The serial site of Quanzhou illustrates the city's vibrancy as a maritime emporium during the Song and Yuan periods (10th - 14th centuries AD) and its interconnection with the Chinese hinterland."),
 "1110": ("International trade / East-West encounter", "It bears witness to one of the earliest and longest-lasting encounters between China and the West, based on the vibrancy of international trade."),
 "1541": ("Sino-foreign exchange", "...this island off the southern coast of the Chinese empire suddenly became an important window for Sino-foreign exchanges."),
 "1112": ("Overseas Chinese / emigration", "They reflect the significant role of émigré Kaiping people in the development of several countries in South Asia, Australasia and North America, during the late 19th and early 20th centuries."),
 "1765": ("Global trade (porcelain)", "...the production of handicraft porcelain from the 10th to 19th centuries that had profound industrial, artistic and trading influences in China and beyond globally."),
 "1389": ("Cross-cultural (Mongol-Han)", "...the site was a unique attempt to assimilate the nomadic Mongolian and Han Chinese cultures."),
}

TL = [("105","Dongzhai Port Nature Reserve","1996"),("106","The Alligator Sinensis Nature Reserve","1996"),("107","Poyang Nature Reserve","1996"),("108","The Lijiang River Scenic Zone at Guilin","1996"),("1622","Yalong, Tibet","2001"),("1623","Yangtze Gorges Scenic Spot","2001"),("1624","Jinfushan Scenic Spot","2001"),("1625","Heaven Pit and Ground Seam Scenic Spot","2001"),("1626","Hua Shan Scenic Area","2001"),("1627","Yandang Mountain","2001"),("1629","Nanxi River","2001"),("1631","Maijishan Scenic Spots","2001"),("1632","Wudalianchi Scenic Spots","2001"),("1633","Haitan Scenic Spots","2001"),("1634","Dali Chanshan Mountain and Erhai Lake Scenic Spot","2001"),("5320","Sites for Liquor Making in China","2008"),("5322","Ancient Residences in Shanxi and Shaanxi Provinces","2008"),("5324","City Walls of the Ming and Qing Dynasties","2008"),("5327","Slender West Lake and Historic Urban Area in Yangzhou","2008"),("5328","The Ancient Waterfront Towns in the South of Yangtze River","2008"),("5335","Chinese Section of the Silk Road: Land routes in Henan Province, Shaanxi Province, Gansu Province, Qinghai Province, Ningxia Hui Autonomous Region, and Xinjiang Uygur Autonomous Region; Sea Routes in Ningbo City, Zhejiang Province and Quanzhou City, Fujian Province - from Western-Han Dynasty to Qing Dynasty","2008"),("5337","Fenghuang Ancient City","2008"),("5338","Site of Southern Yue State","2008"),("5341","Baiheliang Ancient Hydrological Inscription","2008"),("5344","Miao Nationality Villages in Southeast Guizhou Province: The villages of Miao Nationality at the Foot of Leigong Mountain in Miao Ling Mountains","2008"),("5347","Karez Wells","2008"),("5351","Expansion Project of Imperial Tombs of the Ming and Qing Dynasties: King Lujian's Tombs","2008"),("5353","The Four Sacred Mountains as an Extension of Mt. Taishan","2008"),("5532","Taklimakan Desert—Populus euphratica Forests","2010"),("5533","China Altay","2010"),("5535","Karakorum-Pamir","2010"),("5803","Wooden Structures of Liao Dynasty—Wooden Pagoda of Yingxian County, Main Hall of Fengguo Monastery of Yixian County","2013"),("5804","Sites of Hongshan Culture: The Niuheliang Archaeological Site, the Hongshanhou Archaeological Site, and Weijiawopu Archaeological Site","2013"),("5806","Ancient Porcelain Kiln Site in China","2013"),("5808","SanFangQiXiang","2013"),("5813","Dong Villages","2013"),("5814","Lingqu Canal","2013"),("5815","Diaolou Buildings and Villages for Tibetan and Qiang Ethnic Groups","2013"),("5816","Archaeological Sites of the Ancient Shu State: Site at Jinsha and Joint Tombs of Boat-shaped Coffins in Chengdu City, Sichuan Province; Site of Sanxingdui in Guanghan City, Sichuan Province 29C.BC-5C.BC","2013"),("5989","Xinjiang Yardang","2015"),("5990","Dunhuang Yardangs","2015"),("5992","Tianzhushan","2015"),("5993","Jinggangshan--North Wuyishan (Extension of Mount Wuyi)","2015"),("5994","ShuDao","2015"),("5995","Tulin-Guge Scenic and Historic Interest Areas","2015"),("6093","The Chinese Section of the Silk Roads","2016"),("6184","Guancen Mountain -- Luya Mountain","2017"),("6185","Hulun Buir Landscape & Birthplace of Ancient Minority","2017"),("6186","Qinghai Lake","2017"),("6187","Scenic and historic area of Sacred Mountains and Lakes","2017"),("6188","Taihang Mountain","2017"),("6190","Vertical Vegetation Landscape and Volcanic Landscape in Changbai Mountain","2017"),("6380","Huangguoshu Scenic Area","2019"),("6381","Guizhou Triassic Fossil Sites","2019"),("6582","Hainan Tropical Rainforest and the Traditional Settlement of Li Ethnic Group","2022"),("6619","Fujian Minjiang River Estuary: The ecotone between marine and terrestrial biogeographical regions","2022"),("6768","Hezheng Mammalian Fossil Sites of Gansu","2024"),("6840","Guizhou Shuanghe Cave","2025"),("6841","Zhangye Colorful Hills","2025"),("7021","Jehol Biota Fossil Sites in Chaoyang","2026")]
TL_NOTES = {
 "5335": ("Silk Roads + Maritime", "2008 entry: land routes in six provinces plus sea routes at Ningbo and Quanzhou. Earliest TL entry to put the sea routes under the Silk Road name."),
 "6093": ("Silk Roads + Maritime", "Submitted 22/02/2016 by China's National Commission for UNESCO; criteria (i)-(vi). Serial, China only: 32 land-route sites (6 provinces) + 31 sea-route sites (Fujian, Zhejiang, Guangdong, Jiangsu). No partner countries named. Describes routes 'connecting Asia. Africa and Europe'."),
}

# Sites outside China whose UNESCO text links them to China / Chinese / Silk Roads
# (id, name, states, year, transnational, where found, link type, quote)
OTHER = [
 ("1223","Melaka and George Town, Historic Cities of the Straits of Malacca","Malaysia","2008","0","Full OUV statement on site page only (not in feed text)","Chinese community / trade","Criterion (ii): Melaka and George Town represent exceptional examples of multi-cultural trading towns in East and Southeast Asia, forged from the mercantile and exchanges of Malay, Chinese, and Indian cultures and three successive European colonial powers for almost 500 years..."),
 ("948","Hoi An Ancient Town","Viet Nam","1999","0","Full OUV statement on site page only (not in feed text)","Chinese community / trade","The town reflects a fusion of indigenous and foreign cultures (principally Chinese and Japanese with later European influences) that combined to produce this unique survival."),
 ("1610","Ombilin Coal Mining Heritage of Sawahlunto","Indonesia","2019","0","Short description (feed)","Chinese labour / migration","The workforce was recruited from the local Minangkabau people and supplemented by Javanese and Chinese contract workers, and convict labourers from Dutch-controlled areas."),
 ("502","Historic City of Vigan","Philippines","1999","0","Short description (feed)","Chinese community / trade","Its architecture reflects the coming together of cultural elements from elsewhere in the Philippines, from China and from Europe..."),
 ("677","Baroque Churches of the Philippines","Philippines","1993","0","Short description (feed)","Chinese craftsmen","Their unique architectural style is a reinterpretation of European Baroque by Chinese and Philippine craftsmen."),
 ("1328","Central Sector of the Imperial Citadel of Thang Long - Hanoi","Viet Nam","2010","0","Short description (feed)","Chinese political presence / influence","It was constructed on the remains of a Chinese fortress dating from the 7th century... at the crossroads between influences coming from China in the north and the ancient Kingdom of Champa in the south."),
 ("1091","Complex of Koguryo Tombs","Democratic People's Republic of Korea","2004","0","Short description (feed)","Shared history (Koguryo)","...the Koguryo Kingdom, one of the strongest kingdoms in nowadays northeast China and half of the Korean peninsula... Only about 90 out of more than 10,000 Koguryo tombs discovered in China and Korea so far, have wall paintings."),
 ("1537","Ancient City of Qalhat","Oman","2018","0","Short description (feed)","Maritime trade","The Ancient City bears unique archaeological testimony to the trade links between the east coast of Arabia, East Africa, India, China and South-East Asia."),
 ("144","Ruins of Kilwa Kisiwani and Ruins of Songo Mnara","United Republic of Tanzania","1981","0","Short description (feed)","Maritime trade (porcelain)","From the 13th to the 16th century, the merchants of Kilwa dealt in gold, silver, pearls, perfumes, Arabian crockery, Persian earthenware and Chinese porcelain."),
 ("365","Khami Ruins National Monument","Zimbabwe","1986","0","Short description (feed)","Trade goods","The discovery of objects from Europe and China shows that Khami was a major centre for trade over a long period of time."),
 ("1246","Iwami Ginzan Silver Mine and its Cultural Landscape","Japan","2007","0","Short description (feed)","Maritime trade","...port towns from where it was shipped to Korea and China."),
 ("972","Gusuku Sites and Related Properties of the Kingdom of Ryukyu","Japan","2000","0","Justification of criteria (feed)","Maritime trade / exchange","For several centuries the Ryukyu islands served as a centre of economic and cultural interchange between south-east Asia, China, Korea, and Japan..."),
 ("660","Buddhist Monuments in the Horyu-ji Area","Japan","1993","0","Short description (feed)","Cultural influence","...they illustrate the adaptation of Chinese Buddhist architecture and layout to Japanese culture... the introduction of Buddhism to Japan from China by way of the Korean peninsula."),
 ("688","Historic Monuments of Ancient Kyoto (Kyoto, Uji and Otsu Cities)","Japan","1994","0","Short description (feed)","Cultural influence","Built in A.D. 794 on the model of the capitals of ancient China, Kyoto was the imperial capital of Japan..."),
 ("870","Historic Monuments of Ancient Nara","Japan","1998","0","Justification of criteria (feed)","Cultural influence","...the evolution of Japanese architecture and art as a result of cultural links with China and Korea..."),
 ("1142","Sacred Sites and Pilgrimage Routes in the Kii Mountain Range","Japan","2004","0","Short description (feed)","Cultural influence","...Buddhism, which was introduced from China and the Korean Peninsula."),
 ("1757","Ancient Capitals of Asuka and Fujiwara","Japan","2026","0","Short description (feed)","Cultural influence","...a centralised state system modelled after the Chinese codified law known as ritsuryō."),
 ("1477","Baekje Historic Areas","Republic of Korea","2015","0","Short description (feed)","Cultural influence / exchange","...exchanges between the ancient East Asian kingdoms in Korea, China and Japan."),
 ("1498","Seowon, Korean Neo-Confucian Academies","Republic of Korea","2019","0","Short description (feed)","Cultural influence","The seowons illustrate a historical process in which Neo-Confucianism from China was adapted to Korean conditions."),
 ("1439","Namhansanseong","Republic of Korea","2014","0","Short description (feed)","Conflict / influence","...rebuilt... in anticipation of an attack from the Sino-Manchu Qing dynasty... based on Chinese and Japanese influences..."),
 ("559","Royal Domain of Drottningholm","Sweden","1991","0","Short description (feed)","Chinoiserie (European taste)","With its castle, perfectly preserved theatre (built in 1766), Chinese pavilion and gardens..."),
 ("1759","The Cemetery Complexes of the Xiongnu Nobility","Mongolia","2026","0","Short description (feed)","Silk Roads / Han interaction","They interacted with the Han Chinese and controlled nodes along the main parts of the Steppe Route of the Silk Roads."),
 ("1675","Silk Roads: Zarafshan-Karakum Corridor","Tajikistan, Turkmenistan, Uzbekistan","2023","1","Site name + short description (feed)","Silk Roads (no China component)","The Zarafshan-Karakum Corridor is a key section of the Silk Roads in Central Asia that connects other corridors from all directions."),
 ("1346","Tabriz Historic Bazaar Complex","Iran (Islamic Republic of)","2010","0","Short description (feed)","Silk Roads","...one of the most important commercial centres on the Silk Road."),
 ("1544","Historic City of Yazd","Iran (Islamic Republic of)","2017","0","Short description (feed)","Silk Roads","...close to the Spice and Silk Roads."),
 ("1230","Sulaiman-Too Sacred Mountain","Kyrgyzstan","2009","0","Short description (feed)","Silk Roads","...at the crossroads of important routes on the Central Asian Silk Roads."),
 ("1627","Cultural Heritage Sites of Ancient Khuttal","Tajikistan","2025","0","Short description (feed)","Silk Roads","...reflecting its role from the 7th to 16th centuries in Silk Roads trade."),
 ("1518","Archaeological Site of Ani","Türkiye","2016","0","Short description (feed)","Silk Roads","...profited from control of one branch of the Silk Road."),
 ("208","Cultural Landscape and Archaeological Remains of the Bamiyan Valley","Afghanistan","2003","0","Justification of criteria (feed)","Silk Roads","...an important Buddhist centre on the Silk Road..."),
 ("1448","Landscapes of Dauria","Mongolia, Russian Federation","2017","1","Short description (feed)","Geographic mention only","...the Daurian Steppe eco-region, which extends from eastern Mongolia into Russian Siberia and northeastern China."),
 ("340","Keoladeo National Park","India","1985","0","Short description (feed)","Geographic mention only","...wintering areas for large numbers of aquatic birds from Afghanistan, Turkmenistan, China and Siberia."),
 ("1012","Kinabalu Park","Malaysia","2000","0","Short description (feed)","Geographic mention only","...flora from the Himalayas, China, Australia, Malaysia..."),
 ("1574","Amami-Oshima Island, Tokunoshima Island, Northern part of Okinawa Island, and Iriomote Island","Japan","2021","0","Short description (feed)","Geographic mention only","...the serial site forms an arc on the boundary of the East China Sea and Philippine Sea..."),
]

OTHER_REC = [
 ("Intangible heritage (joint)","Ong Chun/Wangchuan/Wangkang ceremony, rituals and related practices for maintaining the sustainable connection between man and the ocean","China + Malaysia","2020","Representative List; decision 15.COM 8.B.22","Diaspora / Maritime","Joint China-Malaysia nomination. UNESCO text: 'Developed in China's Minnan region between the fifteenth and seventeenth centuries, the element is now centered in the coastal areas of Xiamen Bay and Quanzhou Bay, as well as in the Chinese communities in Melaka, Malaysia.'","https://ich.unesco.org/en/RL/ong-chun-wangchuan-wangkang-ceremony-rituals-and-related-practices-for-maintaining-the-sustainable-connection-between-man-and-the-ocean-01608"),
 ("Intangible heritage (joint)","Urtiin Duu, traditional folk long song","Mongolia + China","2005 proclaimed; 2008 inscribed","Representative List","Shared cross-border tradition","Joint Mongolia-China element.","https://ich.unesco.org/en/RL/urtiin-duu-traditional-folk-long-song-00115"),
 ("Intangible heritage (parallel national)","Mongolian art of singing, Khoomei (China) vs Mongolian traditional art of Khöömei (Mongolia)","China (2009); Mongolia (2010)","2009 / 2010","Representative List, two separate national inscriptions","Framing contest","Same tradition inscribed separately by each state a year apart. Useful case for 'who defines heritage'.","https://ich.unesco.org/en/RL/mongolian-art-of-singing-khoomei-00210 ; https://ich.unesco.org/en/RL/mongolian-traditional-art-of-khoomei-00396"),
 ("Intangible heritage (China only)","Mazu belief and customs","China","2009","Representative List","Diaspora (lead)","China-only inscription of a practice with temples across the Chinese diaspora. Lead for the historical-presence layer; not yet checked against UNESCO text.","https://ich.unesco.org/en/state/china-CN?info=elements-on-the-lists"),
 ("Intangible heritage (China only)","Nanyin","China","2009","Representative List","Diaspora (lead)","China-only inscription (Quanzhou music) also practised in Southeast Asian Chinese communities. Lead only; not yet checked against UNESCO text.","https://ich.unesco.org/en/state/china-CN?info=elements-on-the-lists"),
 ("Intangible heritage (China only)","Watertight-bulkhead technology of Chinese junks","China","2010","Urgent Safeguarding List","Maritime (lead)","Maritime technology; possible link to Maritime Silk Road framing. Lead only.","https://ich.unesco.org/en/state/china-CN?info=elements-on-the-lists"),
 ("Memory of the World","Qiaopi and Yinxin Correspondence and Remittance Documents from Overseas Chinese","China","2013 (international); 2012 Asia-Pacific register (Teochew qiaopi)","International Register","Diaspora","Letters and remittances between overseas Chinese and home villages. Directly relevant to overseas Chinese heritage. UNESCO page could not be fetched (robots); details via secondary source.","https://www.unesco.org/en/memory-world/qiaopi-and-yinxin-correspondence-and-remittance-documents-overseas-chinese"),
 ("UNESCO programme","UNESCO Silk Roads Project / Programme (incl. 1990-91 Maritime Route Expedition)","Multinational","1988-1998 project; programme ongoing","Social and Human Sciences Sector","Silk Roads / Maritime","Sen (2023) links the 1990-91 expedition ('90 scholars from 26 countries', 21 ports) to the rise of the 'Maritime Silk Road' concept in the PRC. Key for tracking framing over time.","https://en.unesco.org/silkroad"),
 ("UNESCO capacity-building project","Safeguarding World Heritage along the Silk Roads (Silk Roads transnational nomination support)","Kazakhstan, Kyrgyzstan, Tajikistan, Turkmenistan, Uzbekistan, China and others","2011-2014 (Phase I); 2015-2018 (Phase II)","UNESCO/Japan Funds-in-Trust, run by World Heritage Centre","Silk Roads","Supported the corridor nominations, including Chang'an-Tianshan (2014). Shows Japan, not China, as funder of this track.","https://southsouth-galaxy.org/solution/safeguarding-world-heritage-along-the-silk-roads/"),
 ("UNESCO category 2 centre (in China)","World Heritage Institute of Training and Research for the Asia and the Pacific Region (WHITRAP)","China (serves Asia-Pacific)","Established under UNESCO auspices (year to confirm)","Category 2 centre; Beijing, Shanghai, Suzhou","Chinese institutional engagement","Chinese-hosted UNESCO training institute serving other States Parties. Source for training and partnership data.","https://whitrap.pku.edu.cn/About_us/Introduction_Center.htm"),
]

DOCS_1442 = [
 ("Nomination file 1442","Nomination","2014","https://whc.unesco.org/uploads/nominations/1442.pdf"),
 ("Advisory Body Evaluation (ICOMOS)","Evaluation","2014","https://whc.unesco.org/document/152690"),
 ("Maps of inscribed property","Maps","2014","https://whc.unesco.org/document/132728"),
 ("Decision 38COM 8B.24 (inscription)","Decision","2014","https://whc.unesco.org/en/decisions/6110"),
 ("Decision 40COM 7B.34","Decision (state of conservation)","2016","https://whc.unesco.org/en/decisions/6699"),
 ("Decision 41COM 7B.88","Decision (state of conservation)","2017","https://whc.unesco.org/en/decisions/7087"),
 ("Decision 42COM 7B.5","Decision (state of conservation)","2018","https://whc.unesco.org/en/decisions/7234"),
 ("Decision 44COM 7B.22","Decision (state of conservation)","2021","https://whc.unesco.org/en/decisions/7740"),
 ("Decision 45COM 7B.156","Decision (state of conservation)","2023","https://whc.unesco.org/en/decisions/8160"),
 ("Decision 47COM 7B.67","Decision (state of conservation)","2025","https://whc.unesco.org/en/decisions/8791"),
 ("ICOMOS Advisory Mission Report - Talgar component (Kazakhstan)","Mission report","2016","https://whc.unesco.org/document/142401"),
 ("Joint WHC/ICOMOS Reactive Monitoring Mission Report (Kazakhstan)","Mission report","2016","https://whc.unesco.org/document/158586"),
 ("Periodic Reporting Cycle 3, Section II","Periodic report","2023","https://whc.unesco.org/document/217919"),
 ("State of conservation report (multilingual)","SOC report","2024","https://whc.unesco.org/document/218280"),
 ("State of conservation report","SOC report","2020","https://whc.unesco.org/document/180694"),
 ("State of conservation report (Kazakhstan)","SOC report","2018","https://whc.unesco.org/document/166338"),
 ("State of conservation report (Kazakhstan)","SOC report","2017","https://whc.unesco.org/document/156862"),
 ("State of conservation report (China)","SOC report","2017","https://whc.unesco.org/document/165239"),
 ("State of conservation report (Kazakhstan)","SOC report","2016","https://whc.unesco.org/document/139846"),
 ("State of conservation report (Kyrgyzstan)","SOC report","2016","https://whc.unesco.org/document/139897"),
 ("Summary state of conservation report (China)","SOC report","2015","https://whc.unesco.org/document/139854"),
]

ACCESS = [
 ("World Heritage List - XML feed (EN)","https://whc.unesco.org/en/list/xml/","Bulk download","Yes: 1,273 properties. Fields: id_number, site, states, iso_code, category, criteria_txt, date_inscribed, secondary_dates, danger, transnational, regions, location, short_description, justification, geolocations (one point per component).","Best base layer. id_number is a stable key. Short text only: no full OUV statement."),
 ("World Heritage List - XML feed (ZH)","https://whc.unesco.org/zh/list/xml/","Bulk download","Yes: same 1,273 records in Chinese.","Gives UNESCO's own Chinese names and descriptions for comparing framing. Gaps: e.g. 1442 Silk Roads, 1111 Hani Terraces, 1223 Melaka, 1541 Kulangsu have empty Chinese names."),
 ("World Heritage List - XLS/XLSX, KML, GeoRSS, RSS","https://whc.unesco.org/en/syndication","Bulk download","Yes.","XLSX and KML are handy for mapping. Same content as XML."),
 ("Site pages (/en/list/{id}/)","https://whc.unesco.org/en/list/1223/","Web page","Partly: full Statement of OUV, brief synthesis, criteria text.","Holds the richer wording (e.g. Chinese communities in Melaka, Hoi An) that the feed lacks. Not syndicated: UNESCO terms say non-syndicated sections may not be scraped. Read manually or ask permission."),
 ("Documents pages (/en/list/{id}/documents/)","https://whc.unesco.org/en/list/1442/documents/","Web page + PDFs","Yes, per site: nomination file, ICOMOS/IUCN evaluation, decisions, maps, mission, SOC and periodic reports.","Key for dated framings (nomination vs evaluation vs decisions). PDFs are large (1442 nomination file ~1 GB)."),
 ("Tentative Lists (/en/tentativelists/{id}/)","https://whc.unesco.org/en/statesparties/cn","Web page","Partly: no bulk feed found.","China has 60 TL entries. Other countries' TL entries mentioning China cannot be searched in bulk without scraping; use manual search or request data."),
 ("Decisions database","https://whc.unesco.org/en/decisions/","Web page","Searchable by keyword and session.","Use for UNESCO's official wording over time (e.g. 'Maritime Silk Road' in decisions)."),
 ("Intangible Cultural Heritage lists","https://ich.unesco.org/en/state/china-CN?info=elements-on-the-lists","Web page","Partly: list per state; multinational status only on element pages.","China: 45 elements. Joint ones found: Ong Chun/Wangchuan/Wangkang (with Malaysia), Urtiin Duu (with Mongolia)."),
 ("Memory of the World","https://www.unesco.org/en/memory-world","Web page","Partly.","unesco.org blocked automated fetching (robots.txt). Read manually."),
 ("Wikidata (crosswalk)","https://www.wikidata.org","SPARQL API","Yes: links WH IDs to multilingual names and coordinates.","Not reachable from this workspace; test in Phase 2."),
]


# ---------- build ----------
F = "Arial"
HDR = PatternFill("solid", fgColor="1F3A5F")
HFONT = Font(name=F, bold=True, color="FFFFFF")
BODY = Font(name=F, size=10)
LINK = Font(name=F, size=10, color="0563C1", underline="single")
FILL = PatternFill("solid", fgColor="FFF2CC")

wb = Workbook()

def sheet(title, headers, rows, widths, link_cols=(), fill_cols=()):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font, cell.fill = HFONT, HDR
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for r in rows:
        ws.append(list(r))
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = BODY
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if cell.column in link_cols and isinstance(cell.value, str) and cell.value.startswith("http"):
                cell.hyperlink = cell.value.split(" ; ")[0]
                cell.font = LINK
            if cell.column in fill_cols:
                cell.fill = FILL
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 32
    return ws

# 1 China WH list
wh = json.load(open("data/raw/unesco_whc_feed_china_extract_2026-09-30.json", encoding="utf-8"))
rows = []
for r in sorted(wh, key=lambda x: (x[4], int(x[0]))):
    i, name, states, cat, yr, sec, tn, npoi, lat, lon, crit = r
    others = ", ".join(s for s in states.split(",") if s != "China")
    flag, quote = CN_FLAGS.get(i, ("", ""))
    rows.append([int(i), name, ZH.get(i, ""), cat, crit, int(yr), sec,
                 "Yes" if tn == "1" else "No", others, npoi, float(lat), float(lon),
                 flag, quote, f"{WHC}/en/list/{i}/", f"{WHC}/en/list/{i}/documents/",
                 "", ""])
sheet("China WH List",
      ["UNESCO ID", "Site name (EN)", "Site name (ZH, UNESCO feed)", "Category", "Criteria",
       "Year inscribed", "Extensions", "Transnational", "Other States Parties",
       "Component points", "Latitude", "Longitude", "Outward link in UNESCO text",
       "UNESCO wording (EN short description)", "Site page", "Documents page",
       "Your notes", "Verification status"],
      rows, [9, 42, 26, 10, 16, 9, 11, 11, 20, 10, 11, 11, 22, 60, 30, 34, 26, 14],
      link_cols=(15, 16), fill_cols=(17, 18))

# 2 China Tentative List
rows = []
for i, name, yr in TL:
    flag, note = TL_NOTES.get(i, ("", ""))
    rows.append([int(i), name, int(yr), flag, note, f"{WHC}/en/tentativelists/{i}/", "", ""])
sheet("China Tentative List",
      ["UNESCO TL ID", "Name (as submitted)", "Year submitted", "Silk Road / maritime link",
       "Notes", "UNESCO page", "Your notes", "Verification status"],
      rows, [10, 60, 10, 20, 60, 36, 26, 14], link_cols=(6,), fill_cols=(7, 8))

# 3 Other WH sites linked to China
rows = [[int(i), n, ZH.get(i, ""), s, int(y), "Yes" if t == "1" else "No", where, lt, q,
         f"{WHC}/en/list/{i}/", ""] for (i, n, s, y, t, where, lt, q) in OTHER]
sheet("Sites Abroad Linked to China",
      ["UNESCO ID", "Site name (EN)", "Site name (ZH, UNESCO feed)", "States Parties",
       "Year inscribed", "Transnational", "Where UNESCO mentions it", "Link type (my coding)",
       "UNESCO wording (verbatim excerpt)", "Site page", "Your notes"],
      rows, [9, 40, 24, 22, 9, 11, 26, 24, 70, 30, 26], link_cols=(10,), fill_cols=(11,))

# 4 Other UNESCO records
sheet("Other UNESCO Records",
      ["Record type", "Name", "States", "Year(s)", "List / programme", "Relevance to project",
       "Notes", "Link"],
      OTHER_REC, [22, 44, 24, 18, 26, 20, 70, 40], link_cols=(8,))

# 5 Documents for 1442
sheet("Docs - Silk Roads 1442",
      ["Document", "Type", "Year", "URL"], DOCS_1442, [58, 26, 8, 50], link_cols=(4,))

# 6 Data access
sheet("Data Access",
      ["Source", "URL", "Access method", "Systematic / structured?", "Notes for the project"],
      ACCESS, [34, 42, 16, 50, 70], link_cols=(2,))

# README with count formulas
ws = wb["Sheet"]
ws.title = "README"
lines = [
 ("Mapping Global China - Phase 1a: UNESCO data inventory", None),
 (f"As of {ASOF}. Built from UNESCO's World Heritage syndication feeds (EN and ZH), UNESCO site, Tentative List and ICH pages.", None),
 ("", None),
 ("Summary", None),
 ("China World Heritage properties", "=COUNTA('China WH List'!A2:A200)"),
 ("  of which Cultural", "=COUNTIF('China WH List'!D2:D200,\"Cultural\")"),
 ("  of which Natural", "=COUNTIF('China WH List'!D2:D200,\"Natural\")"),
 ("  of which Mixed", "=COUNTIF('China WH List'!D2:D200,\"Mixed\")"),
 ("  of which transnational", "=COUNTIF('China WH List'!H2:H200,\"Yes\")"),
 ("  with an outward link in UNESCO's text", "=COUNTIF('China WH List'!M2:M200,\"?*\")"),
 ("China Tentative List entries", "=COUNTA('China Tentative List'!A2:A200)"),
 ("  with a Silk Road / maritime link", "=COUNTIF('China Tentative List'!D2:D200,\"?*\")"),
 ("World Heritage sites abroad linked to China in UNESCO text", "=COUNTA('Sites Abroad Linked to China'!A2:A200)"),
 ("  excluding geographic-only mentions", "=COUNTA('Sites Abroad Linked to China'!A2:A200)-COUNTIF('Sites Abroad Linked to China'!H2:H200,\"Geographic mention only\")"),
 ("", None),
 ("Sheets", None),
 ("China WH List", "All 61 inscribed properties with Chinese names, coordinates (first component point), transnational flag, and any outward link stated in UNESCO's own description."),
 ("China Tentative List", "All 60 entries. Two put the maritime routes under the Silk Roads name (2008, 2016)."),
 ("Sites Abroad Linked to China", "Properties in other countries whose UNESCO text mentions China, Chinese people or the Silk Roads. Link type is my coding and should be reviewed."),
 ("Other UNESCO Records", "Intangible heritage, Memory of the World, Silk Roads programme and centres relevant to the project."),
 ("Docs - Silk Roads 1442", "Full document trail for the one transnational property China is part of."),
 ("Data Access", "What each UNESCO source offers, how to get it, gaps and terms of use."),
 ("", None),
 ("Legend", None),
 ("Yellow columns", "For you to fill in: notes and verification status."),
 ("Coordinates", "First component point from UNESCO's feed; serial sites have several points (see 'Component points')."),
 ("", None),
 ("Terms of use", "UNESCO: 'Any republication, online or in any other form, of any UNESCO/WHC data requires prior written authorization.' Required credit: Copyright © 1992-2026 UNESCO/World Heritage Centre. All rights reserved. Fine for internal research; ask UNESCO before publishing a map layer built on this data."),
]
for i, (a, b) in enumerate(lines, 1):
    ws.cell(row=i, column=1, value=a).font = Font(name=F, size=10)
    if b is not None:
        c = ws.cell(row=i, column=2, value=b)
        c.font = Font(name=F, size=10)
        c.alignment = Alignment(wrap_text=True, vertical="top")
for r in (1,):
    ws.cell(row=r, column=1).font = Font(name=F, size=14, bold=True)
for r in (4, 16, 24):
    ws.cell(row=r, column=1).font = Font(name=F, size=11, bold=True)
ws.column_dimensions["A"].width = 52
ws.column_dimensions["B"].width = 100

out = "data/processed/Mapping_Global_China_Phase1a_UNESCO.xlsx"
wb.save(out)
print("saved", out)
