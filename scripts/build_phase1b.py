"""Add Phase 1b (Chinese sources) sheets to the Phase 1a workbook.

Run from the repo root after build_phase1a.py: python scripts/build_phase1b.py
All values were collected on 2026-09-30 from the URLs recorded in each row.
"""
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

SRC = "data/processed/Mapping_Global_China_Phase1a_UNESCO.xlsx"
OUT = "data/processed/Mapping_Global_China_Phase1_Sources.xlsx"

F = "Arial"
HDR = PatternFill("solid", fgColor="1F3A5F")
HFONT = Font(name=F, bold=True, color="FFFFFF")
BODY = Font(name=F, size=10)
LINK = Font(name=F, size=10, color="0563C1", underline="single")
FILL = PatternFill("solid", fgColor="FFF2CC")

wb = load_workbook(SRC)

def sheet(title, headers, rows, widths, link_cols=(), fill_cols=()):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font, cell.fill = HFONT, HDR
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for r in rows:
        ws.append(list(r))
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = BODY
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if cell.column in link_cols and isinstance(cell.value, str) and cell.value.startswith("http"):
                cell.hyperlink = cell.value
                cell.font = LINK
            if cell.column in fill_cols:
                cell.fill = FILL
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 32
    return ws

# ---------------- 1b source inventory ----------------
# Source | Owner type | URL | Coverage (layers/regions) | Language | How systematic | Access method (preliminary) | Historical versions | Framing notes | Reliability | Your notes
INV = [
 ("National Cultural Heritage Administration (国家文物局, NCHA)", "Chinese government (heritage regulator)", "http://www.ncha.gov.cn/",
  "Institutional engagement (aid restoration, joint archaeology, Asian Cultural Heritage Protection Action, ACHA); UNESCO and nominations (China's own and transnational). Global, with focus on Asia, Silk Roads, Africa.",
  "Chinese", "Semi-structured. No dedicated international column: overseas items are scattered across 文物要闻, 工作动态 and 政务公开. Article URLs are dated (/art/YYYY/M/D/).",
  "Scraping via dated article URLs. Site search (/jsearch/) did not respond to plain URL queries; needs browser automation. Loads from outside China.",
  "Yes: archive back to at least 2014; five-year plans 2017, 2021, 2026 form a dated series.",
  "Official policy language: 文明交流互鉴, 一带一路, 服务国家外交大局, 全球文明倡议, 中国智慧和方案.", "High for what China says and does; not neutral.", ""),
 ("NCHA five-year plans (十三五 2017, 十四五 2021, 十五五 2026)", "Chinese government (policy documents)", "http://www.ncha.gov.cn/art/2026/7/29/art_2318_48456.html",
  "All four layers at policy level: named aid projects, joint nominations (Maritime Silk Road, Tea Road), joint archaeology target regions.",
  "Chinese", "Highly systematic: same document type every five years, comparable sections.",
  "Manual (3 documents). 14th plan is hosted on ndrc.gov.cn.",
  "Yes: 2017, 2021, 2026. Best single source for tracking framing change (Phase 5).",
  "2017: Maritime Silk Road nomination as domestic work. 2026: 'transnational joint nomination' preparatory work; 'serve the overall diplomatic agenda'.", "High (authoritative statement of intent).", ""),
 ("Ministry of Culture and Tourism (文化和旅游部, MCT)", "Chinese government (culture ministry)", "https://www.mct.gov.cn/",
  "Intangible heritage policy and UNESCO ICH nominations (incl. joint China-Malaysia element); cultural exchange; China Cultural Centres abroad.",
  "Chinese (some English)", "Semi-structured news and policy columns; plans issued as PDF attachments.",
  "Scraping of news pages; PDFs for plans (ICH 15th five-year plan issued 2026-09-21 as PDF, not yet read).",
  "Yes: news archive; five-year ICH plans.", "Not yet sampled in depth.", "High for official position.", ""),
 ("China Intangible Cultural Heritage Network (中国非物质文化遗产网, ihchina.cn)", "Chinese government-affiliated (run by China ICH Protection Center under MCT)", "https://www.ihchina.cn/project.html",
  "National ICH list: 3,610 items; 45 UNESCO-inscribed items. Descriptions often mention overseas Chinese, Southeast Asia, Taiwan.",
  "Chinese", "Highly structured: item ID, number (e.g. Ⅹ-85), title, category, new/extension, year and batch, province/city code, description.",
  "Open JSON endpoint: /getProject.html?keywords=...&category_id=16&limit=10&p=1 (tested; returns JSON).",
  "Partial: listing year and batch per item; descriptions not dated.",
  "Diaspora framing: Nanyin as 'spiritual bond' for overseas compatriots and Taiwan compatriots; Mazu spread 'following the footsteps of overseas Chinese'. National title 闽台送王船 (Fujian-Taiwan) vs UNESCO's China + Malaysia.", "High for official framing.", ""),
 ("Chinese embassy websites (e.g. Cambodia, Malaysia, Kenya, Uzbekistan)", "Chinese government (diplomatic)", "https://kh.china-embassy.gov.cn/",
  "Institutional engagement and narrative layers per host country: heritage exhibitions, restoration projects, ambassador essays on shared history, overseas-Chinese affairs (领侨动态).",
  "Chinese (most also English)", "Semi-structured. Common Foreign Ministry template on every *.china-embassy.gov.cn site; dated URLs (/column/YYYYMM/tYYYYMMDD_id.htm). Some list pages render by script.",
  "Scraping by dated URL pattern; search runs through a central Foreign Ministry engine (fmprc.gov.cn/irs-c-web/search.shtml) filterable by site. Some old URLs (pre-2013) now redirect to the homepage: use the Wayback Machine.",
  "Yes: Kenya archive back to 2005, Malaysia to 2012 in samples.",
  "Heavy framing: 'civilisations learning from each other' (互学互鉴), Zheng He as friendship, 'Silk Road spirit', 'community with a shared future'.", "High for official narrative; low for facts on sites.", ""),
 ("Chinese Academy of Cultural Heritage (中国文化遗产研究院, CACH)", "Chinese government research institute (under NCHA)", "https://www.chinanews.com.cn/gj/2022/08-29/9839659.shtml",
  "Institutional engagement: lead agency for aid restoration (Angkor Chau Say Tevoda 1998, Ta Keo 2010, Royal Palace 2018).",
  "Chinese", "Own site not yet tested; project information mostly via NCHA and media interviews.", "To test in Phase 2.", "Yes: project history since 1990s.",
  "Frames Angkor work as China being among the earliest participants in UNESCO's international safeguarding campaign.", "High.", ""),
 ("CASS Institute of Archaeology (中国社会科学院考古研究所) and cssn.cn", "Chinese state research institute", "https://www.cssn.cn/skgz/bwyc/202605/t20260522_5992725.shtml",
  "Joint archaeology abroad: Uzbekistan (2012-), Honduras Copan (2015-), Egypt Monthu Temple (2018-), Romania (2019-), Greece (2026). 2026 catalogue 《大道东风》 of overseas results.",
  "Chinese", "Semi-structured: news and features on cssn.cn; excavation reports as books.", "Scraping of cssn.cn articles; reports need library access.",
  "Yes: projects dated from 2012.", "'Independent knowledge system' (自主知识体系); from 'archaeological power' to 'archaeological strength'.", "High for project facts.", ""),
 ("Provincial and university archaeology teams (e.g. Shaanxi Institute of Archaeology, Northwest University, Peking University)", "Chinese research institutes / universities", "https://www.chinanews.com.cn/sh/2026/09-24/10703295.shtml",
  "Joint archaeology in Central Asia (Kyrgyzstan Red River site, Shaanxi) and East Africa (Kenya Lamu/Malindi, Peking University, c. 2010-2013).",
  "Chinese, some English", "Scattered: press releases, university news, academic papers.", "Manual; media search.", "Yes, via news archives.",
  "Kenya work framed around Zheng He's voyages and 'rewriting East African coast history'.", "Medium-high.", ""),
 ("Alliance for Cultural Heritage in Asia (亚洲文化遗产保护联盟, ACHA)", "China-led multilateral body", "https://english.www.gov.cn/news/202304/25/content_WS6447673dc6d03ffcca6ec9ee.html",
  "Institutional engagement across Asia: 10 founding states (China, Armenia, Cambodia, DPRK, Iran, Kyrgyzstan, Pakistan, Syria, UAE, Yemen); roadmap includes Silk Roads archaeology, World Heritage nominations, restoration in Türkiye and Syria. NCHA's 2026 plan adds an 'Asian cultural heritage database'.",
  "Chinese, English", "No dedicated website found; information via NCHA and state media.", "Manual.", "From 2021 (dialogue) and 2023 (Xi'an conference).",
  "'Dialogue among civilizations'; Global Civilization Initiative.", "High for membership and plans.", ""),
 ("Belt and Road portal (中国一带一路网, yidaiyilu.gov.cn)", "Chinese government (aggregator)", "https://www.yidaiyilu.gov.cn/",
  "Aggregates heritage and archaeology news framed under Belt and Road.", "Chinese, English", "Semi-structured news aggregator.", "Scraping; not yet tested.", "Yes, since 2017.",
  "Belt and Road framing by definition.", "Medium (secondary, republished).", ""),
 ("Overseas-Chinese institutions (Overseas Chinese History Museum of China; Jinan University Academy of Overseas Chinese Studies; China Federation of Returned Overseas Chinese)", "Chinese state museum / university / mass organisation", "https://hqhryj.jnu.edu.cn/5535/list.htm",
  "Historical Chinese presence layer: diaspora history, qiaopi, associations, temples abroad.",
  "Chinese", "Weakly structured online. Jinan has a documentation centre and journals (《东南亚研究》, 《海外华人研究》) but no open database found.",
  "Manual; library access.", "Journal archives.", "Not yet sampled.", "Medium-high (scholarly).", ""),
 ("CNKI (中国知网) academic literature", "Commercial database (Chinese scholarship)", "https://www.cnki.net/",
  "Chinese academic writing on overseas heritage, Maritime Silk Road, joint archaeology; useful for dating when terms emerge.",
  "Chinese", "Highly structured bibliographic records with dates.", "Subscription: check Columbia library access.", "Yes: decades of dated records.",
  "Good for counting term frequency over time (e.g. 海上丝绸之路 since the 1980s).", "High as a record of discourse.", ""),
]

sheet("1b Chinese Sources",
      ["Source", "Owner type", "URL (example page)", "Coverage (layers, regions)", "Language",
       "How systematic", "Access method (preliminary, for Phase 2)", "Historical versions",
       "Framing notes", "Reliability", "Your notes"],
      INV, [34, 22, 40, 55, 12, 42, 48, 30, 52, 20, 24], link_cols=(3,), fill_cols=(11,))

# ---------------- framing examples ----------------
# Date | Actor | Actor type | Topic / site | Quote (original) | English gloss (my translation) | Framing tags | Source URL
FR = [
 ("2017-02-21", "NCHA (13th five-year plan)", "Government", "Maritime Silk Road nomination",
  "推进花山岩画文化景观、鼓浪屿·历史国际社区、古泉州（刺桐）史迹、良渚遗址、海上丝绸之路保护与申遗工作，推动陆上丝绸之路其他廊道申遗。",
  "Advance protection and nomination of ... Ancient Quanzhou (Zayton), Liangzhu, the Maritime Silk Road; push nomination of other land Silk Road corridors.",
  "Maritime Silk Road; nomination (domestic)", "http://www.ncha.gov.cn/art/2017/2/27/art_2237_43663.html"),
 ("2017-02-21", "NCHA (13th five-year plan)", "Government", "Belt and Road heritage corridor",
  "增进与“一带一路”沿线国家及文化遗产国际组织的交流合作，建设“一带一路”文化遗产长廊。",
  "Increase exchange with Belt and Road countries and international heritage bodies; build a 'Belt and Road cultural heritage corridor'.",
  "Belt and Road", "http://www.ncha.gov.cn/art/2017/2/27/art_2237_43663.html"),
 ("2021-12-01", "NCHA (14th five-year plan)", "Government", "Aid restoration: Angkor, Nepal, Myanmar",
  "推进柬埔寨吴哥古迹、尼泊尔九层神庙、缅甸他冰瑜佛塔等文物保护修复合作项目，持续开展中外联合考古。",
  "Advance restoration cooperation at Angkor (Cambodia), the Nine-Storey Temple (Nepal) and Thatbyinnyu Temple (Myanmar); continue joint archaeology.",
  "Aid restoration; Belt and Road", "https://www.ndrc.gov.cn/fggz/fzzlgh/gjjzxgh/202112/t20211201_1306596.html"),
 ("2021-10-25", "NCHA (Asian Cultural Heritage Protection Action series)", "Government", "China-Cambodia history",
  "新的世纪，官方合作再次恢复，朝着构建人类命运共同体的新目标前进。",
  "In the new century, official cooperation has resumed, moving toward the new goal of building a community with a shared future for mankind.",
  "Shared history; community of shared future", "http://www.ncha.gov.cn/art/2021/10/25/art_722_171543.html"),
 ("2022-02-18", "NCHA news feature", "Government", "Aid restoration history (from Mongolia 1961 to Angkor)",
  "中国文物保护工作者持续为国际濒危文化遗产保护提供中国智慧和方案，让受损的不同文明的代表性建筑焕发光彩。",
  "Chinese heritage professionals keep providing 'Chinese wisdom and solutions' for endangered heritage worldwide, restoring landmark buildings of different civilisations.",
  "Chinese wisdom and solutions", "http://www.ncha.gov.cn/art/2022/2/18/art_722_173029.html"),
 ("2026-07-20", "NCHA (15th five-year plan)", "Government", "Transnational joint nominations",
  "规范跨省域联合申遗项目管理，持续推进海上丝绸之路、万里茶道、中埃水文遗产保护与跨国联合申遗前期工作。",
  "Continue preparatory work for protection and transnational joint nomination of the Maritime Silk Road, the Tea Road, and China-Egypt hydrological heritage.",
  "Maritime Silk Road; nomination (transnational)", "http://www.ncha.gov.cn/art/2026/7/29/art_2318_48456.html"),
 ("2026-07-20", "NCHA (15th five-year plan)", "Government", "Heritage diplomacy",
  "拓展文物交流合作服务国家外交大局，立足文物资源优势，强化机制建设和项目带动，促进文明交流互鉴。",
  "Expand heritage exchange to serve the country's overall diplomatic agenda ... promote exchange and mutual learning among civilisations.",
  "Diplomacy; civilisational exchange", "http://www.ncha.gov.cn/art/2026/7/29/art_2318_48456.html"),
 ("2026-07-20", "NCHA (15th five-year plan)", "Government", "Aid restoration: Angkor, Nepal, Myanmar",
  "健全文物援外机制，推进柬埔寨吴哥古迹、尼泊尔努瓦科特杜巴广场王宫、缅甸他冰瑜佛塔等文物古迹保护合作。",
  "Improve the heritage-aid mechanism; advance cooperation at Angkor, the Nuwakot Durbar Square palace (Nepal) and Thatbyinnyu Temple (Myanmar).",
  "Aid restoration", "http://www.ncha.gov.cn/art/2026/7/29/art_2318_48456.html"),
 ("2026-07-20", "NCHA (15th five-year plan)", "Government", "Joint archaeology target regions",
  "开展南岛语族主要分布地区、环喜马拉雅地区、丝绸之路和海上丝绸之路沿线地区等联合考古，培育5—10个示范项目。",
  "Carry out joint archaeology in Austronesian areas, the Himalayas, and along the Silk Roads and Maritime Silk Road; develop 5-10 demonstration projects.",
  "Joint archaeology; Silk Roads; Maritime Silk Road", "http://www.ncha.gov.cn/art/2026/7/29/art_2318_48456.html"),
 ("2026-07-20", "NCHA (15th five-year plan)", "Government", "Asian heritage database",
  "推动制定亚洲文化遗产保护准则与技术规范，建立亚洲文化遗产数据库与联合实验室。",
  "Promote Asian heritage conservation standards; build an Asian cultural heritage database and joint laboratories.",
  "Standards; data (possible future source)", "http://www.ncha.gov.cn/art/2026/7/29/art_2318_48456.html"),
 ("2026-09-24", "NCHA (joint archaeology results release)", "Government", "Egypt, Tunisia, Kyrgyzstan joint projects",
  "深入践行全球文明倡议，不断完善中外联合考古工作管理制度，积极支持中国考古机构、专家学者“走出去”。",
  "Put the Global Civilization Initiative into practice ... actively support Chinese archaeological institutions and scholars to 'go global'.",
  "Global Civilization Initiative; going global", "https://www.chinanews.com.cn/sh/2026/09-24/10703295.shtml"),
 ("2026-05-22", "CASS (cssn.cn feature)", "Research institute", "Overseas archaeology programme",
  "中国考古不仅在走向世界，更在以自主知识体系深度介入不同文明的交流互鉴。",
  "Chinese archaeology is not only going global but engaging deeply in exchange between civilisations through its own independent knowledge system.",
  "Independent knowledge system", "https://www.cssn.cn/skgz/bwyc/202605/t20260522_5992725.shtml"),
 ("2026-09-08", "Chinese Embassy in Cambodia (Amb. Wang Wenbin)", "Embassy", "China-Cambodia heritage exhibition; Angkor",
  "中华文明与高棉文明互学互鉴，相映生辉，谱写了世界文明美美与共的佳话。",
  "Chinese and Khmer civilisations have learned from each other and shone together, a model story of world civilisations' shared beauty.",
  "Civilisational exchange", "https://kh.china-embassy.gov.cn/dssghd/202609/t20260908_12018158.htm"),
 ("2016-10-28", "Chinese Embassy in Malaysia (ambassador's essay)", "Embassy", "Melaka; Zheng He",
  "明朝郑和七下西洋，其中5次驻节马六甲，在加深两国人民友谊的同时，更有力维护了周边地区近百年的和平与繁荣。",
  "Zheng He's seven voyages included five stays in Melaka, deepening friendship and maintaining nearly a century of regional peace and prosperity.",
  "Zheng He; peace narrative; Maritime Silk Road", "https://my.china-embassy.gov.cn/sgxx/dsjh/201611/t20161105_1790289.htm"),
 ("2016-10-28", "Chinese Embassy in Malaysia (ambassador's essay)", "Embassy", "Chinese community in Malaysia",
  "近代以来，大批华人南下定居马来西亚，经世代交替，始终与当地人民和谐相处、共存共荣，华人文化传统已与马来西亚历史发展进程深深融合。",
  "Large numbers of Chinese settled in Malaysia in modern times ... Chinese cultural traditions have deeply merged with Malaysia's historical development.",
  "Diaspora; integration", "https://my.china-embassy.gov.cn/sgxx/dsjh/201611/t20161105_1790289.htm"),
 ("2024-02-23", "Chinese Embassy in Kenya (Amb. Zhou Pingjian)", "Embassy", "Lamu, Malindi, Mombasa; Zheng He",
  "肯尼亚是中华文明与非洲文明最早交汇地之一，拉穆、马林迪、蒙巴萨等滨海地区至今流传着郑和船队的友好佳话。",
  "Kenya is one of the earliest meeting places of Chinese and African civilisations; coastal Lamu, Malindi and Mombasa still tell stories of Zheng He's fleet.",
  "Zheng He; Silk Road spirit", "https://ke.china-embassy.gov.cn/sgxx/dszs_145423/202402/t20240223_11249135.htm"),
 ("2006 (listed)", "ihchina national ICH list: Nanyin (Quanzhou)", "Government-affiliated database", "Nanyin in Southeast Asia",
  "泉州南音还流播到菲律宾、印尼、新加坡、马来西亚、泰国、缅甸、越南等国家，成为维系海外侨胞和台湾同胞乡情的精神纽带，对增进民族认同感也起到了积极作用。",
  "Nanyin spread to the Philippines, Indonesia, Singapore, Malaysia, Thailand, Myanmar and Vietnam, becoming a spiritual bond for overseas compatriots and Taiwan compatriots, and strengthening national identity.",
  "Diaspora; national identity", "https://www.ihchina.cn/project.html"),
 ("2006 (listed)", "ihchina national ICH list: Mazu rituals (Putian)", "Government-affiliated database", "Mazu temples worldwide",
  "妈祖文化逐步传播至我国的沿江、沿海和台港澳地区，并随着华侨华人的脚印逐步传播到世界上的五大洲二十多个国家。",
  "Mazu culture spread along China's rivers and coasts and to Taiwan, Hong Kong and Macao, and followed the footsteps of overseas Chinese to over 20 countries on five continents.",
  "Diaspora", "https://www.ihchina.cn/project.html"),
 ("2011 (listed)", "ihchina national ICH list: Mazu rituals (Dongtou)", "Government-affiliated database", "Mazu and cross-Strait relations",
  "妈祖信仰是海峡两岸关系重要的桥梁和纽带，是大陆和台湾文化交融的平台，对于促进祖国统一具有积极意义。",
  "Mazu belief is an important bridge for cross-Strait relations ... and positive for promoting national reunification.",
  "Cross-Strait; reunification", "https://www.ihchina.cn/project.html"),
 ("2011 (listed)", "ihchina national ICH list: 民间信俗（闽台送王船）", "Government-affiliated database", "Wangchuan ceremony (compare UNESCO 2020)",
  "“王爷”信仰广泛流传于闽南沿海及台湾渔村，尤其盛于南台湾，与中台湾的妈祖信仰并称。",
  "Wangye belief is widespread in southern Fujian coastal areas and Taiwanese fishing villages. National title frames it as Fujian-Taiwan; UNESCO's 2020 joint inscription frames it as China + Malaysia (Melaka).",
  "Cross-Strait vs diaspora framing", "https://www.ihchina.cn/project.html"),
]

sheet("1b Framing Examples",
      ["Date", "Actor", "Actor type", "Topic / site", "Original wording (verbatim)",
       "English gloss (my translation)", "Framing tags (my coding)", "Source"],
      FR, [13, 34, 18, 30, 60, 60, 26, 44], link_cols=(8,))

# ---------------- Chinese projects abroad ----------------
PR = [
 ("Mongolia", "Xingren Temple and Bogd Khan Palace (lama temples), Ulaanbaatar", "Aid restoration", "1959-1961", "Chinese experts Yu Mingqian and Li Zhujun", "Described by NCHA as the start of China's overseas heritage work.", "http://www.ncha.gov.cn/art/2022/2/18/art_722_173029.html"),
 ("Cambodia", "Angkor: Chau Say Tevoda", "Aid restoration", "1998 (ten-year project)", "CACH", "Within UNESCO's International Coordinating Committee for Angkor.", "http://www.ncha.gov.cn/art/2022/2/18/art_722_173029.html"),
 ("Cambodia", "Angkor: Ta Keo", "Aid restoration", "2010 (eight-year project)", "CACH", "", "http://www.ncha.gov.cn/art/2022/2/18/art_722_173029.html"),
 ("Cambodia", "Angkor: Royal Palace site", "Aid restoration", "2018 (signed)", "CACH", "Named again in the 2026 plan.", "http://www.ncha.gov.cn/art/2022/2/18/art_722_173029.html"),
 ("Nepal", "Nine-Storey Temple (Basantapur), Kathmandu", "Aid restoration", "Named in 2021 plan", "Not stated in sources read", "Kathmandu Valley World Heritage site.", "https://www.ndrc.gov.cn/fggz/fzzlgh/gjjzxgh/202112/t20211201_1306596.html"),
 ("Nepal", "Nuwakot Durbar Square palace", "Aid restoration", "Named in 2026 plan", "Not stated in sources read", "", "http://www.ncha.gov.cn/art/2026/7/29/art_2318_48456.html"),
 ("Myanmar", "Thatbyinnyu Temple, Bagan", "Aid restoration", "Named in 2021 and 2026 plans", "Not stated in sources read", "Bagan World Heritage site.", "http://www.ncha.gov.cn/art/2026/7/29/art_2318_48456.html"),
 ("Uzbekistan", "Fergana Basin sites (Mingtepe, Monzatepe, Eratan)", "Joint archaeology", "2012-", "CASS Institute of Archaeology", "Silk Roads region.", "https://www.cssn.cn/skgz/bwyc/202605/t20260522_5992725.shtml"),
 ("Honduras", "Copan Maya site", "Joint archaeology", "2015-", "CASS Institute of Archaeology", "Non-Chinese-heritage site: comparison case.", "https://www.cssn.cn/skgz/bwyc/202605/t20260522_5992725.shtml"),
 ("Egypt", "Monthu Temple, Karnak (Luxor)", "Joint archaeology", "2018-", "CASS Institute of Archaeology", "Results released 2026-09-24.", "https://www.chinanews.com.cn/sh/2026/09-24/10703295.shtml"),
 ("Romania", "Site reported as 'Dobrujavatz' (name to confirm)", "Joint archaeology", "2019-", "CASS Institute of Archaeology", "", "https://www.cssn.cn/skgz/bwyc/202605/t20260522_5992725.shtml"),
 ("Greece", "Angelokastro site (as reported)", "Joint archaeology", "2026", "CASS Institute of Archaeology", "", "https://www.cssn.cn/skgz/bwyc/202605/t20260522_5992725.shtml"),
 ("Tunisia", "Ben Arous forest site", "Joint archaeology", "Results 2026", "NCHA Archaeological Research Center", "Carthaginian-Roman transition.", "https://www.chinanews.com.cn/sh/2026/09-24/10703295.shtml"),
 ("Kyrgyzstan", "Red River site, Chu valley", "Joint archaeology", "Results 2026", "Shaanxi Academy of Archaeology", "Silk Roads region (near the Chang'an-Tianshan corridor).", "https://www.chinanews.com.cn/sh/2026/09-24/10703295.shtml"),
 ("Kenya", "Lamu archipelago and Malindi", "Joint archaeology", "c. 2010-2013 (reported)", "Peking University (per media reports)", "Framed around Zheng He's voyages; secondary sources only so far.", "https://ke.china-embassy.gov.cn/sgxx/dszs_145423/202402/t20240223_11249135.htm"),
]
sheet("1b Chinese Projects Abroad",
      ["Country", "Site", "Type", "Years", "Chinese institution", "Notes", "Source", "Verification status"],
      PR, [14, 40, 18, 22, 34, 44, 44, 16], link_cols=(7,), fill_cols=(8,))

# ---------------- README additions ----------------
ws = wb["README"]
ws["A1"] = "Mapping Global China - Phase 1 source inventory (1a UNESCO, 1b Chinese sources)"
r = ws.max_row + 2
add = [
 ("Phase 1b (added 2026-09-30)", None),
 ("Chinese sources surveyed", "=COUNTA('1b Chinese Sources'!A2:A200)"),
 ("Dated framing examples", "=COUNTA('1b Framing Examples'!A2:A200)"),
 ("Chinese projects abroad (from sources read)", "=COUNTA('1b Chinese Projects Abroad'!A2:A200)"),
 ("1b Chinese Sources", "Inventory in the roadmap's 1e format: coverage, language, how systematic, access, historical versions, framing, reliability."),
 ("1b Framing Examples", "Dated verbatim quotes from government, embassy, research and ICH sources, with English gloss. Seed data for the narrative layer and Phase 5."),
 ("1b Chinese Projects Abroad", "Aid restoration and joint archaeology projects named in the sources read. Seed for the institutional-engagement layer."),
]
for i, (a, b) in enumerate(add):
    ws.cell(row=r + i, column=1, value=a).font = Font(name=F, size=11 if i == 0 else 10, bold=(i == 0))
    if b is not None:
        c = ws.cell(row=r + i, column=2, value=b)
        c.font = Font(name=F, size=10)
        c.alignment = Alignment(wrap_text=True, vertical="top")

wb.save(OUT)
print("saved", OUT, len(INV), len(FR), len(PR))
